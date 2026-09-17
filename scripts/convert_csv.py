"""
convert_csv.py

Converts avante_n_cn7_profile_614_all_pids.csv + avante_n_cn7_profile_614_all_pids_cn.xlsx
into data/pids.js for use in the WeChat mini-program.

Run with:
    D:\miniconda3\envs\obd_stats\python.exe scripts/convert_csv.py

Output:
    data/pids.js
"""

import csv
import json
import re
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).parent.parent
CSV_PATH = ROOT / "data" / "avante_n_cn7_profile_614_all_pids.csv"
XLSX_PATH = ROOT / "data" / "avante_n_cn7_profile_614_all_pids_cn.xlsx"
OUT_PATH = ROOT / "data" / "pids.js"


def load_chinese_names(xlsx_path: Path) -> dict[str, str]:
    """Load parameter_id -> Chinese name from the new full-data xlsx.

    The new file has a single sheet with header row containing
    'parameter_id' and 'parameter_name_cn' columns.
    """
    wb = openpyxl.load_workbook(xlsx_path)
    ws = wb.worksheets[0]

    # Locate columns by header name (robust against column order changes)
    header = [c.value for c in ws[1]]
    try:
        id_col = header.index("parameter_id")
        cn_col = header.index("parameter_name_cn")
    except ValueError as e:
        raise RuntimeError(f"Expected column not found in xlsx: {e}") from e

    id_to_cn: dict[str, str] = {}
    for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True)):
        param_id = str(row[id_col] or "").strip()
        cn_name = str(row[cn_col] or "").strip()
        if param_id:
            id_to_cn[param_id] = cn_name
    return id_to_cn


def parse_float_or_none(val: str) -> float | None:
    """Parse a numeric string, return None if empty or non-numeric."""
    val = val.strip()
    if not val:
        return None
    try:
        return float(val)
    except ValueError:
        return None


def parse_int_or_none(val: str) -> int | None:
    val = val.strip()
    if not val:
        return None
    try:
        return int(val)
    except ValueError:
        return None


def load_csv_rows(csv_path: Path) -> list[dict]:
    """
    Load CSV using Python's csv module to correctly handle fields
    that may contain commas (e.g., conv_rule).
    Returns a list of row dicts using the header row as keys.
    """
    with open(csv_path, encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        return list(reader)


def build_pid_entry(
    row: dict,
    cn_name: str,
) -> dict | None:
    """
    Build a compact PID entry dict from a CSV row + Chinese name.
    Returns None if the row is malformed (e.g., non-integer data_pos).
    """
    param_id = row.get("parameter_id", "").strip()
    name_en = row.get("parameter_name_en", "").strip()
    ecu = row.get("ecu_name_en", "").strip()

    request_cmd_raw = row.get("request_cmd", "").strip()
    # Strip spaces from request_cmd for compact storage (e.g., "22 E0 01" -> "22E001")
    request_cmd = request_cmd_raw.replace(" ", "").upper()

    request_header = row.get("request_header", "").strip().upper()
    response_header = row.get("response_header", "").strip().upper()
    pid_did = row.get("pid_did", "").strip().replace(" ", "").upper()

    data_pos_str = row.get("data_pos", "").strip()
    data_size_type_str = row.get("data_size_type", "").strip()
    data_size_str = row.get("data_size", "").strip()
    bit_type_str = row.get("bit_type", "").strip()
    conv_rule = row.get("conv_rule", "").strip()

    min_val = parse_float_or_none(row.get("min", ""))
    max_val = parse_float_or_none(row.get("max", ""))
    unit = row.get("unit", "").strip()
    point_str = row.get("point", "").strip()

    session_open_cmd = row.get("session_open_cmd", "").strip()
    tester_present_cmd = row.get("tester_present_cmd", "").strip()

    # Validate required numeric fields
    data_pos = parse_int_or_none(data_pos_str)
    data_size_type = parse_int_or_none(data_size_type_str)
    data_size = parse_int_or_none(data_size_str)
    bit_type = parse_int_or_none(bit_type_str)
    point = parse_int_or_none(point_str)

    if data_pos is None or data_size_type is None or data_size is None or bit_type is None:
        return None  # skip malformed rows

    # Only accept data_size_type 0 or 1 (skip CSV parsing artifacts)
    if data_size_type not in (0, 1):
        return None

    if not param_id or not name_en or not ecu or not request_cmd:
        return None

    return {
        "id": param_id,
        "name_en": name_en,
        "name_cn": cn_name,
        "ecu": ecu,
        "rh": request_header,
        "rsp": response_header,
        "cmd": request_cmd,
        "did": pid_did,
        "dp": data_pos,
        "dst": data_size_type,
        "ds": data_size,
        "bt": bit_type,
        "cr": conv_rule,
        "mn": min_val,
        "mx": max_val,
        "unit": unit,
        "pt": point if point is not None else 0,
        "soc": session_open_cmd,
        "tpc": tester_present_cmd,
    }


def main() -> None:
    print(f"Loading Chinese names from {XLSX_PATH.name}...")
    id_to_cn = load_chinese_names(XLSX_PATH)
    print(f"  Loaded {len(id_to_cn)} entries.")

    print(f"Loading CSV from {CSV_PATH.name}...")
    rows = load_csv_rows(CSV_PATH)
    print(f"  Loaded {len(rows)} rows.")

    pids: list[dict] = []
    skipped_malformed = 0
    skipped_dst = 0
    missing_cn = 0

    for row in rows:
        param_id = row.get("parameter_id", "").strip()
        cn_name = id_to_cn.get(param_id, "")
        if not cn_name:
            missing_cn += 1
            # Fall back to English name if no Chinese translation
            cn_name = row.get("parameter_name_en", "").strip()

        # Check data_size_type before full build for stats
        dst_str = row.get("data_size_type", "").strip()
        dst_val = parse_int_or_none(dst_str)
        if dst_val not in (0, 1):
            skipped_dst += 1
            continue

        entry = build_pid_entry(row, cn_name)
        if entry is None:
            skipped_malformed += 1
            continue
        pids.append(entry)

    print(f"\nResults:")
    print(f"  Valid entries:        {len(pids)}")
    print(f"  Skipped (bad dst):    {skipped_dst}")
    print(f"  Skipped (malformed):  {skipped_malformed}")
    print(f"  Missing Chinese name: {missing_cn}")

    # Stats by data_size_type
    dst0 = sum(1 for p in pids if p["dst"] == 0)
    dst1 = sum(1 for p in pids if p["dst"] == 1)
    print(f"\n  data_size_type=0 (numeric): {dst0}")
    print(f"  data_size_type=1 (bit):     {dst1}")

    # ECU distribution
    from collections import Counter
    ecu_counts = Counter(p["ecu"] for p in pids)
    print(f"\n  ECU distribution:")
    for ecu, cnt in ecu_counts.most_common():
        print(f"    {cnt:4d}  {ecu}")

    # Write output as a JS module (pids.js) so WeChat mini-program require() works.
    # JSON files cannot be require()'d directly in the mini-program runtime.
    print(f"\nWriting {OUT_PATH} ...")
    json_str = json.dumps(pids, ensure_ascii=False, separators=(",", ":"))
    with open(OUT_PATH, "w", encoding="utf-8") as f:
        f.write("module.exports=")
        f.write(json_str)
        f.write("\n")

    size_kb = OUT_PATH.stat().st_size / 1024
    print(f"Done. File size: {size_kb:.1f} KB ({len(pids)} entries)")

    if size_kb > 1800:
        print("WARNING: File may exceed WeChat mini-program package limits!")


if __name__ == "__main__":
    main()
